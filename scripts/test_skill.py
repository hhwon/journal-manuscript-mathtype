#!/usr/bin/env python3
"""Offline regression tests. Run native MathType smoke tests separately online."""
from __future__ import annotations
import json
from pathlib import Path
import sys
import tempfile
import unittest
from zipfile import ZipFile
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

from crossrefs import make_anchors, validate_refs
from build_paper import make_doc
from patch_equation_bookmarks import patch
from validate_docx import inspect
from lint_manuscript import lint
from safe_text_patch import patch_text

ROOT=Path(__file__).resolve().parent.parent


class ManuscriptSkillTest(unittest.TestCase):
    def test_anchors_forward_and_backward_refs(self):
        spec={'sections':[{'blocks':[
            {'type':'paragraph','segments':[{'text':'See '},{'ref':'eq:one'}]},
            {'type':'equation','latex':'x=1','id':'eq:one'},
            {'type':'figure','caption':'A','id':'fig:a'}]}]}
        by_id,by_block=make_anchors(spec)
        validate_refs(spec,by_id)
        self.assertEqual(by_id['eq:one'].number,1)
        self.assertEqual(by_id['fig:a'].kind,'figure')
    def test_duplicate_ids_fail(self):
        spec={'sections':[{'blocks':[{'type':'equation','id':'a'},
                                     {'type':'equation','id':'a'}]}]}
        with self.assertRaisesRegex(ValueError,'Duplicate'):
            make_anchors(spec)
    def test_missing_ref_fails(self):
        spec={'sections':[{'blocks':[{'type':'paragraph',
                                      'segments':[{'ref':'eq:missing'}]}]}]}
        refs,_=make_anchors(spec)
        with self.assertRaisesRegex(ValueError,'unknown reference'):
            validate_refs(spec,refs)
    def test_nonsequential_number_fails(self):
        spec={'sections':[{'blocks':[{'type':'equation','number':3}]}]}
        with self.assertRaisesRegex(ValueError,'Non-sequential'):
            make_anchors(spec)
    def test_native_math_backend_fails_closed_without_dependency(self):
        from build_paper import backend
        try:
            deps,_=backend(False,True)
        except RuntimeError as err:
            self.assertIn('MathType',str(err))
        else:
            self.assertIsNotNone(deps)
    def test_preview_crossrefs_and_figure_link(self):
        spec_path=ROOT/'examples/demo_paper.json'
        spec=json.loads(spec_path.read_text(encoding='utf-8'))
        with tempfile.TemporaryDirectory() as tmp:
            output=Path(tmp)/'preview.docx'
            result=make_doc(spec,output,True,None,False,spec_path.parent)
            self.assertTrue(output.is_file())
            self.assertTrue(result['layout_preview_only'])
            audit=inspect(output)
            n_refs=sum(1 for sec in spec['sections'] for b in sec['blocks'] for sg in b.get('segments',[]) if isinstance(sg,dict) and 'ref' in sg)
            self.assertEqual(audit['word_ref_fields'],n_refs)
            self.assertIn('NON-PRODUCTION',audit['warnings'][0])
            with ZipFile(output) as z:
                x=z.read('word/document.xml').decode('utf-8')
                self.assertIn('SEQ Figure',x)
                self.assertIn('SEQ Table',x)
                self.assertIn('MS_EQ_',x)
                self.assertIn('MS_FIG_',x)
                self.assertIn('MS_TAB_',x)
                self.assertNotIn('Equation Chapter',x)
    def test_submission_blocks_synthetic_demo(self):
        demo=json.loads((ROOT/'examples/demo_paper.json').read_text(encoding='utf-8'))
        issues=lint(demo,submission=True)
        self.assertTrue(any(row['level']=='error' and row['location']=='sample_data' for row in issues))

    def test_plain_no_math_has_clean_build(self):
        spec={'title':'Plain paper','sections':[{'heading':'1. Background',
             'blocks':[{'type':'paragraph','text':'The model is defined.'}]}]}
        with tempfile.TemporaryDirectory() as tmp:
            output=Path(tmp)/'paper.docx'
            report=make_doc(spec,output,False,None,False,Path(tmp))
            self.assertTrue(report['structural_qa'])
            self.assertEqual(inspect(output)['ole_bins'],0)
    def test_safe_text_patch_preserves_all_other_docx_parts(self):
        with tempfile.TemporaryDirectory() as tmp:
            src=Path(tmp)/'old.docx'
            out=Path(tmp)/'new.docx'
            doc=Document()
            doc.add_paragraph('The preliminary results are described.')
            doc.save(src)
            result=patch_text(src,out,[{'old':'preliminary','new':'numerical'}])
            self.assertTrue(result['ole_binaries_preserved_byte_for_byte'])
            self.assertEqual(Document(out).paragraphs[0].text,
                             'The numerical results are described.')
            self.assertEqual(Document(src).paragraphs[0].text,
                             'The preliminary results are described.')
            with ZipFile(src) as a, ZipFile(out) as b:
                for name in a.namelist():
                    if name != 'word/document.xml':
                        self.assertEqual(a.read(name),b.read(name))

    def test_safe_patch_rejects_word_field_results(self):
        from crossrefs import _field
        with tempfile.TemporaryDirectory() as tmp:
            src=Path(tmp)/'old.docx'
            out=Path(tmp)/'new.docx'
            doc=Document()
            p=doc.add_paragraph('Figure ')
            _field(p,r'SEQ Figure \* ARABIC','1')
            doc.save(src)
            with self.assertRaisesRegex(ValueError,'expected 1 plain text occurrences; got 0'):
                patch_text(src,out,[{'old':'1','new':'9'}])
            self.assertFalse(out.exists())

    def test_patch_equation_bookmark(self):
        spec={'sections':[{'blocks':[{'type':'equation','id':'eq:test'}]}]}
        anchors,_=make_anchors(spec)
        anchor=anchors['eq:test']
        with tempfile.TemporaryDirectory() as tmp:
            path=Path(tmp)/'artificial.docx'
            doc=Document()
            p=doc.add_paragraph()
            ppr=p._p.get_or_add_pPr()
            tabs=OxmlElement('w:tabs');ppr.append(tabs)
            tab=OxmlElement('w:tab');tab.set(qn('w:val'),'center');tabs.append(tab)
            r=p.add_run(); t=OxmlElement('w:tab');r._r.append(t)
            r=p.add_run(); e=etree.Element('{urn:schemas-microsoft-com:office:office}OLEObject');r._r.append(e)
            r=p.add_run(); t=OxmlElement('w:tab');r._r.append(t)
            p.add_run('(1)')
            doc.save(path)
            self.assertEqual(patch(path,[anchor]),1)
            with ZipFile(path) as z:
                root=etree.fromstring(z.read('word/document.xml'))
                names=root.xpath('.//*[local-name()="bookmarkStart"]/@*[local-name()="name"]')
                self.assertIn(anchor.bookmark,names)

try:
    import docx_equation  # noqa: F401
    HAVE_BACKEND = True
except ImportError:
    HAVE_BACKEND = False


class HouseStyleAndGenuineMathTypeTest(unittest.TestCase):
    """Rules found with Windows Word + MathType (references/mathtype-facts.md, format-profile.md)."""

    def test_template_indents_and_spacing(self):
        d = Document(ROOT / 'assets' / 'manuscript-template.docx')
        for name in ('Heading 1', 'Heading 2', 'MS Figure', 'MS Figure Caption', 'MS Table Caption',
                     'MS Equation', 'MS Table Cell'):
            self.assertEqual(d.styles[name].paragraph_format.first_line_indent, 0, name)
        for name in ('Normal', 'MS Body', 'MS Abstract'):
            pf = d.styles[name].paragraph_format
            self.assertEqual((pf.space_before, pf.space_after), (0, 0), name)

    def test_latex_mathml(self):
        from latex_mathml import to_mathml
        m = to_mathml(r'\mathbf{x};\;\boldsymbol\theta')
        self.assertIn('mathvariant="bold">x<', m)
        self.assertIn(';', m)                                   # a plain semicolon stays (\; becomes a space)
        self.assertFalse(any(ord(c) > 0xFFFF for c in m))

    @unittest.skipUnless(HAVE_BACKEND, 'docx-equation not installed')
    def test_directory_tree_repair(self):
        import io, olefile
        from docx_equation.mathtype.ole import build_mathtype_ole_object
        import cfb_treecheck, ole_fix
        blob = build_mathtype_ole_object(b'\x05' + b'\x00' * 20, 'Equation.DSMT4')
        fixed = ole_fix.fix_dir_tree(ole_fix.set_root_clsid(blob))
        self.assertEqual(cfb_treecheck.check(fixed), [])
        a, b = olefile.OleFileIO(io.BytesIO(blob)), olefile.OleFileIO(io.BytesIO(fixed))
        for s in a.listdir():
            self.assertEqual(a.openstream(s).read(), b.openstream(s).read())

    @unittest.skipUnless(HAVE_BACKEND, 'docx-equation not installed')
    def test_profile_and_symbol_positions(self):
        import mtef_profile as P
        P.configure(greek='times', matrix_align='left')
        self.assertEqual(P.encode_char('='), bytes([2, 4, 0x86, 0x3D, 0, 0x3D]))
        self.assertEqual(P.encode_char('∫')[-1], 0xF2)
        self.assertEqual(P.encode_char('θ'), bytes([2, 0, 0x84]) + 'θ'.encode('utf-16-le'))
        self.assertEqual(P.encode_char(' ')[2:], bytes([0x98, 0x02, 0xEF]))
        P.configure(greek='symbol')
        self.assertEqual(P.encode_char('θ')[-1], 0x71)          # Symbol position of theta
        pre = P._preamble('DSMT4')
        self.assertIn(b'Symbol\x00', pre)
        self.assertIn(P.encode_dimensions(P.FACTORY_SPACING), pre)
        P.configure(greek='times', matrix_align='left')

    @unittest.skipUnless(HAVE_BACKEND, 'docx-equation not installed')
    def test_eqp_preferences(self):
        import mtef_profile as P
        from mtef_font_audit import _dimensions
        eqp = ROOT / 'examples' / 'mathtype-factory.eqp'
        P.configure(eqp=str(eqp))
        a = P._preamble('DSMT4')
        P.configure(greek='symbol', matrix_align='center')
        self.assertEqual(a, P._preamble('DSMT4'))               # the factory .eqp = the built-in factory profile
        with tempfile.TemporaryDirectory() as tmp:
            f = Path(tmp) / 'p.eqp'
            f.write_text(eqp.read_text().replace('Full=12 pt', 'Full=11 pt').replace('Script=58 %', 'Script=80 %'))
            prof = P.configure(eqp=str(f))
            self.assertEqual((prof.full_pt, prof.size_pct('Script')), (11.0, 80.0))
            sizes, _ = _dimensions(P._eqn_prefs(), 2)
            self.assertEqual((sizes[0]['value'], sizes[1]['value']), ('11', '80'))
        P.configure(greek='times', matrix_align='left')

    @unittest.skipUnless(HAVE_BACKEND, 'docx-equation not installed')
    def test_preview_transform(self):
        import preview_math as V
        from latex_mathml import to_mathml
        m = V.transform(to_mathml(r'\int_V a\,dV=b_{,i}+\bar P', True))
        self.assertIn('mathsize="150%"', m)                     # MathType integral
        self.assertIn('lspace="0.21em"', m)                     # relation spacing
        self.assertIn('lspace="0em" rspace="0em">,', m.replace('data-sym="1" ', ''))   # no spacing in scripts
        self.assertIn('mathbackground="black"', m)              # bar accent as a rule


if __name__=='__main__':
    unittest.main(verbosity=2)
