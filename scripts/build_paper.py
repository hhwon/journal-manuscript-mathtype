#!/usr/bin/env python3
"""Structured JSON -> journal DOCX with strict MathType-OLE output on Linux.

This program never silently replaces MathType by OMML or an equation picture.
--layout-preview is an explicitly marked, non-production mode.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

from crossrefs import add_caption_prefix, add_reference, make_anchors, validate_refs

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / 'assets' / 'manuscript-template.docx'


def backend(preview: bool, needed: bool):
    if preview or not needed:
        return None, None
    try:
        import docx_equation
        from docx_equation import (EquationRegistry, EquationStyle, ExportOptions,
                                   MathTypeOptions, NumberingOptions)
        from latex2mathml.converter import convert
        import importlib.metadata
        version = importlib.metadata.version('docx-equation')
        if version != '0.3.3':
            raise RuntimeError(
                f'The manuscript adapter is tested against API 0.3.3; detected {version}. '
                'Install the Git SHA pinned in requirements-generator.txt.')
        return (EquationRegistry, EquationStyle, ExportOptions, MathTypeOptions, NumberingOptions), convert
    except ImportError as exc:
        raise RuntimeError(
            'Native MathType conversion requires the pinned docx-equation 0.3.3 '
            'and latex2mathml. Run: python -m pip install -r requirements-generator.txt. '
            f'Missing module: {exc}') from exc


def make_mathml(value: dict, convert, display=False):
    if 'mathml' in value:
        payload = value['mathml']
        if not isinstance(payload, str) or '<math' not in payload:
            raise ValueError('MathML equation must contain a valid <math> root')
        return payload
    if 'latex' in value:
        raw = value['latex']
        if not isinstance(raw, str) or not raw.strip():
            raise ValueError('Equation LaTeX source cannot be empty')
        try:
            from latex_mathml import to_mathml
            return to_mathml(raw, display=display)
        except Exception as exc:
            raise ValueError(f'LaTeX to MathML conversion failed for {raw!r}: {exc}') from exc
    raise ValueError('Equation requires latex or mathml')


def has_math(spec):
    def recurse(obj):
        if isinstance(obj, dict):
            if 'latex' in obj or 'mathml' in obj:
                return True
            return any(recurse(v) for v in obj.values())
        if isinstance(obj, list):
            return any(recurse(x) for x in obj)
        return False
    return recurse(spec)


def segments(value):
    if isinstance(value, str):
        return [{'text': value}]
    if isinstance(value, dict):
        if 'segments' in value:
            return value['segments']
        return [value]
    raise ValueError(f'Text must be a string or an object with segments, not {type(value).__name__}')


def add_segments(p, value, registry, convert, preview, refs):
    for seg in segments(value):
        if not isinstance(seg, dict):
            raise ValueError('Each segment must be a JSON object')
        if 'ref' in seg:
            add_reference(p, refs[seg['ref']], prefix=seg.get('prefix'))
        elif 'latex' in seg or 'mathml' in seg:
            if preview:
                r = p.add_run(' [FORMULA PREVIEW ONLY] ')
                r.italic = True
            else:
                registry.add_inline(p, make_mathml(seg, convert))
        else:
            r = p.add_run(str(seg.get('text', '')))
            if 'bold' in seg:
                r.bold = bool(seg['bold'])
            if 'italic' in seg:
                r.italic = bool(seg['italic'])
            if 'underline' in seg:
                r.underline = bool(seg['underline'])
    return p


def paragraph(doc, value, registry, convert, preview, refs, style='MS Body'):
    p = doc.add_paragraph(style=style)
    add_segments(p, value, registry, convert, preview, refs)
    return p


def _cell_padding(cell, top=75, left=110, bottom=75, right=110):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None:
        mar = OxmlElement('w:tcMar')
        tcPr.append(mar)
    for name, val in (('top', top), ('start', left), ('bottom', bottom), ('end', right)):
        x = mar.find(qn('w:' + name))
        if x is None:
            x = OxmlElement('w:' + name)
            mar.append(x)
        x.set(qn('w:w'), str(val))
        x.set(qn('w:type'), 'dxa')


def three_line(table, header=True):
    """Three-line table: rules at the top, under the header row and at the bottom; no vertical lines."""
    tblPr = table._tbl.tblPr
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for side, val in (('top', 'single'), ('left', 'nil'), ('bottom', 'single'), ('right', 'nil'),
                      ('insideH', 'nil'), ('insideV', 'nil')):
        e = OxmlElement(f'w:{side}')
        e.set(qn('w:val'), val)
        if val == 'single':
            e.set(qn('w:sz'), '8'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), '000000')
        b.append(e)
    tblPr.append(b)
    if header and len(table.rows):
        for c in table.rows[0].cells:
            tb = OxmlElement('w:tcBorders')
            e = OxmlElement('w:bottom')
            e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '6'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), '000000')
            tb.append(e)
            c._tc.get_or_add_tcPr().append(tb)


def table_block(doc, item, anchor, registry, convert, preview, refs):
    cap = doc.add_paragraph(style='MS Table Caption')
    add_caption_prefix(cap, anchor)
    add_segments(cap, item['caption'], registry, convert, preview, refs)
    headings = item.get('headers', [])
    rows = item.get('rows', [])
    if not headings and not rows:
        raise ValueError('Table must have headers or rows')
    cols = len(headings) if headings else len(rows[0])
    if cols < 1:
        raise ValueError('Table requires at least one column')
    if any(len(row) != cols for row in rows):
        raise ValueError(f'Table {anchor.number}: inconsistent column counts')
    table = doc.add_table(rows=1 if headings else 0, cols=cols)
    table.style = 'Table Grid'                     # replaced by three rules below unless item['style'] == 'grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    def put(cell, value, is_heading=False):
        cell.text = ''
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _cell_padding(cell)
        p = cell.paragraphs[0]
        p.style = doc.styles['MS Table Cell']
        add_segments(p, value, registry, convert, preview, refs)
        if is_heading:
            for r in p.runs:
                r.bold = True

    if headings:
        for i, x in enumerate(headings):
            put(table.rows[0].cells[i], x, True)
        trPr = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement('w:tblHeader')
        repeat.set(qn('w:val'), 'true')
        trPr.append(repeat)
        cant = OxmlElement('w:cantSplit')
        trPr.append(cant)
    for row in rows:
        cells = table.add_row().cells
        for i, x in enumerate(row):
            put(cells[i], x)
    if item.get('style', 'three-line') == 'three-line':
        three_line(table, bool(headings))
    if 'note' in item:
        note = paragraph(doc, item['note'], registry, convert, preview, refs, 'MS Note')
        note.paragraph_format.space_before = Pt(4)


def figure_block(doc, item, anchor, registry, convert, preview, refs, base_dir):
    path = Path(item.get('path', ''))
    if not path.is_absolute():
        path = base_dir / path
    if not path.is_file():
        raise FileNotFoundError(f'Figure {anchor.number}: image file not found: {path}')
    if path.suffix.lower() == '.svg':
        import cairosvg
        image_bytes = cairosvg.svg2png(url=str(path), output_width=2200)
        inp = BytesIO(image_bytes)
    else:
        inp = str(path)
    with Image.open(BytesIO(image_bytes) if isinstance(inp, BytesIO) else inp) as image:
        ratio = image.width / max(image.height, 1)
    width = min(6.50, max(.5, float(item.get('width_inches', 5.4))))
    max_height = min(8.0, max(1., float(item.get('max_height_inches', 7.0))))
    width = min(width, max_height * ratio)
    p = doc.add_paragraph(style='MS Figure')
    r = p.add_run()
    shape = r.add_picture(inp, width=Inches(width))
    alt = item.get('alt_text') or item.get('caption', '')
    shape._inline.docPr.set('descr', alt if isinstance(alt, str) else 'Journal manuscript figure')
    cap = doc.add_paragraph(style='MS Figure Caption')
    add_caption_prefix(cap, anchor)
    add_segments(cap, item['caption'], registry, convert, preview, refs)
    cap.paragraph_format.keep_together = True


def make_doc(spec, path, preview, chrome_path, write_manifest, source_dir=None, preview_scale=None,
             profile=None, display_align='left', indent_twips=482):
    if not isinstance(spec, dict) or not isinstance(spec.get('sections'), list):
        raise ValueError('Manuscript root must be an object with a sections array')
    if not spec.get('title'):
        raise ValueError('Manuscript title must not be empty')
    if preview_scale is not None and not (0.10 <= float(preview_scale) <= 1.0):
        raise ValueError('Equation preview scale must be between 0.10 and 1.0 pt/px')
    refs, blocks = make_anchors(spec)
    validate_refs(spec, refs)
    deps, convert = backend(preview, has_math(spec))
    registry = deps[0]() if deps else None
    if deps:
        import mtef_profile                       # genuine-equivalent MTEF (references/mathtype-facts.md)
        mtef_profile.configure(**(profile or {'greek': 'times', 'matrix_align': 'left'}))
        for note in mtef_profile.PROFILE.notes:
            print('equation profile:', note)
        if preview_scale is None:                 # previews: 38 px per em at device scale 2 = the full size
            preview_scale = mtef_profile.PROFILE.full_pt / 76
        mtef_profile.UNMAPPED_SYMBOLS.clear()
        import preview_math                       # MathType-like previews with measured baselines
        preview_math.install()
    doc = Document(TEMPLATE)
    doc.core_properties.title = spec['title']
    # authors: "Name" or ["Name", "1*"] (superscript affiliation marks); affiliations: "text" or ["1", "text"]
    authors = [a if isinstance(a, list) else [a, ''] for a in spec.get('authors', [])]
    if authors:
        doc.core_properties.author = ', '.join(a[0] for a in authors)
    paragraph(doc, spec['title'], registry, convert, preview, refs, 'MS Title')
    if authors:
        pa = doc.add_paragraph(style='MS Authors')
        for k, (name, sup) in enumerate(authors):
            pa.add_run(name)
            if sup:
                pa.add_run(sup).font.superscript = True
            if k < len(authors) - 1:
                pa.add_run(', ')
    for affiliation in spec.get('affiliations', []):
        if isinstance(affiliation, list):
            q = doc.add_paragraph(style='MS Affiliation')
            q.add_run(affiliation[0]).font.superscript = True
            q.add_run(affiliation[1])
        else:
            paragraph(doc, affiliation, registry, convert, preview, refs, 'MS Affiliation')
    if spec.get('corresponding_author'):
        paragraph(doc, 'Corresponding author: ' + spec['corresponding_author'], registry,
                  convert, preview, refs, 'MS Corresponding')
    if spec.get('abstract'):
        doc.add_paragraph('Abstract', style='MS Abstract Heading')
        paragraph(doc, spec['abstract'], registry, convert, preview, refs, 'MS Abstract')
    if spec.get('keywords'):
        p = doc.add_paragraph(style='MS Keywords')
        p.add_run('Keywords: ').bold = True
        kw = '; '.join(spec['keywords'])
        p.add_run(kw if kw.endswith('.') else kw + '.').italic = True          # italic, ending with a full stop
    counts = {'equation': 0, 'figure': 0, 'table': 0}
    eq_anchors = []
    base_dir = Path(source_dir) if source_dir is not None else path.parent
    for section in spec['sections']:
        if section.get('heading'):
            paragraph(doc, section['heading'], registry, convert, preview, refs, 'Heading 1')
        prev_kind = None
        for item in section.get('blocks', []):
            kind = item.get('type')
            if kind == 'paragraph':
                content = {'segments': item['segments']} if 'segments' in item else item.get('text', '')
                p = paragraph(doc, content, registry, convert, preview, refs)
                # a sentence continued after a display equation ("where ...") starts at the margin
                first = (item.get('text') or next((sg.get('text', '') for sg in item.get('segments', []) if isinstance(sg, dict)), '') or '').lstrip()
                if item.get('continues', prev_kind == 'equation' and first[:1].islower()):
                    p.paragraph_format.first_line_indent = Inches(0)
            elif kind in ('subsection', 'heading'):
                style = ('Heading 2' if kind == 'subsection' or item.get('level', 1) == 2
                         else 'Heading 1')
                paragraph(doc, item['heading'], registry, convert, preview, refs, style)
            elif kind == 'equation':
                counts['equation'] += 1
                anchor = blocks[id(item)]
                eq_anchors.append(anchor)
                if preview:
                    p = doc.add_paragraph(style='MS Equation')
                    # Mirror the real engine's center/right tab positions (6.5 in text width).
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p.paragraph_format.left_indent = Inches(0)
                    p.paragraph_format.first_line_indent = Inches(0)
                    p.paragraph_format.tab_stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CENTER)
                    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.50), WD_TAB_ALIGNMENT.RIGHT)
                    p.add_run('\t[FORMULA PREVIEW ONLY]\t')
                    # Preview captions use cached literal text; NEVER release as native MathType.
                    from crossrefs import add_bookmark_start, add_bookmark_end
                    add_bookmark_start(p, anchor)
                    p.add_run(f'({anchor.number})')
                    add_bookmark_end(p, anchor)
                else:
                    registry.add_display(doc, make_mathml(item, convert, display=True), number=anchor.number, style='MS Equation')
            elif kind == 'figure':
                counts['figure'] += 1
                figure_block(doc, item, blocks[id(item)], registry, convert, preview, refs, base_dir)
            elif kind == 'table':
                counts['table'] += 1
                table_block(doc, item, blocks[id(item)], registry, convert, preview, refs)
            elif kind == 'note':
                paragraph(doc, item['text'], registry, convert, preview, refs, 'MS Note')
            else:
                raise ValueError(f'Unsupported block type {kind!r}')
            prev_kind = kind
    if spec.get('acknowledgements'):
        paragraph(doc, 'Acknowledgements', registry, convert, preview, refs, 'Heading 1')
        paragraph(doc, spec['acknowledgements'], registry, convert, preview, refs)
    if spec.get('references'):
        paragraph(doc, 'References', registry, convert, preview, refs, 'Heading 1')
        for k, ref in enumerate(spec['references'], 1):
            # numbered list [n] with a hanging indent (tab stop of the style), unless the entries are already numbered
            numbered = isinstance(ref, str) and not ref.lstrip().startswith('[') and spec.get('reference_style', 'numbered') == 'numbered'
            paragraph(doc, f'[{k}]\t{ref}' if numbered else ref, registry, convert, preview, refs, 'MS References')
    output = Path(path).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    if preview or registry is None or not len(registry):
        doc.save(output)
        result = None
    else:
        _, EquationStyle, ExportOptions, MathTypeOptions, NumberingOptions = deps
        opts = ExportOptions(
            target='mathtype',
            style=EquationStyle(font_family='Times New Roman', east_asia_font='SimSun',
                                font_size_pt=12, number_font_size_pt=12),
            numbering=NumberingOptions(number_format='(1)', use_seq_field=True),
            mathtype=MathTypeOptions(embed_mode='png-preview', mathtype_version='DSMT4',
                                     preview_pt_per_px=preview_scale, preview_font_px=38,
                                     max_width_pt=400, chrome_path=chrome_path),
        )
        try:
            result = registry.save(doc, str(output), opts)
            converted = getattr(result, 'converted', 0)
            if converted != len(registry):
                raise RuntimeError(f'Only {converted}/{len(registry)} equations converted. '
                                   f'Errors: {getattr(result, "errors", [])}')
            from repair_equation_layout import repair
            repair(output)
            from patch_equation_bookmarks import patch
            patch(output, eq_anchors)
            from equation_layout import fix_layout
            fix_layout(output, display_align=display_align, indent_twips=indent_twips)
            import ole_fix
            ole_fix.fix_docx(output)
            from mathtype_gate import gate
            g = gate(output, mtef_profile.UNMAPPED_SYMBOLS)
            if not g['pass']:
                raise RuntimeError('MathType gate failed: ' + json.dumps(g, ensure_ascii=False))
        except Exception:
            output.unlink(missing_ok=True)
            raise
    from validate_docx import inspect
    report = inspect(output, expected=len(registry) if registry else 0,
                     require_mathtype=bool(registry and len(registry)))
    if not report['pass_structural_checks'] and not preview:
        output.unlink(missing_ok=True)
        raise RuntimeError('DOCX QA failure: ' + '; '.join(report['warnings']))
    if write_manifest:
        out_manifest = output.with_suffix('.equations.json')
        payload = {
            'source': str(path),
            'docx_sha256': hashlib.sha256(output.read_bytes()).hexdigest(),
            'equations': [{'id': x.identifier, 'number': x.number,
                           'source': next((item.get('latex') or item.get('mathml'))
                                          for section in spec['sections']
                                          for item in section['blocks']
                                          if item.get('type') == 'equation' and blocks[id(item)] == x)}
                          for x in eq_anchors],
            'windows_mathtype_edit_test': 'not_performed'
        }
        out_manifest.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    return {'output': str(output), **counts, 'inline_and_display_ole_count': report['ole_bins'],
            'layout_preview_only': bool(preview), 'structural_qa': report['pass_structural_checks'],
            'windows_editing_verified': False}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('spec', help='Structured manuscript JSON')
    parser.add_argument('-o', '--output', required=True)
    parser.add_argument('--chrome-path', help='Path to Chromium/Google Chrome for MathML previews')
    parser.add_argument('--layout-preview', action='store_true',
                        help='NON-PRODUCTION: show labelled placeholders instead of MathType OLE')
    parser.add_argument('--math-preview-scale', type=float, default=None,
                        help='Point-per-pixel of the equation previews (default: full size / 76, i.e. 0.158 for 12 pt)')
    # defaults = the house style: Times New Roman only (Greek too), matrices left-aligned, display equations left-aligned
    # at the 0.85 cm body indent with the number right-aligned by a tab. MathType's own defaults:
    # --greek symbol --matrix-align center --display-align center
    parser.add_argument('--greek', choices=['symbol', 'times'], default=None,
                        help='Greek letters in Times New Roman (default without --profile) or in the Symbol font')
    parser.add_argument('--matrix-align', choices=['center', 'left', 'right'], default='left',
                        help='Column alignment inside matrices (default left; MathType factory: center)')
    parser.add_argument('--profile', help='MathType preference file (.eqp: fonts, sizes, spacing), or a JSON file '
                                          '{"eqp": ..., "greek": ..., "matrix_align": ...}')
    parser.add_argument('--display-align', choices=['center', 'left'], default='left',
                        help='Display equations left-aligned at --indent-twips (default) or centred')
    parser.add_argument('--indent-twips', type=int, default=482, help='Left indent of display equations (482 = 0.85 cm)')
    parser.add_argument('--write-equation-manifest', action='store_true',
                        help='Write sidecar JSON mapping display formulas to their source LaTeX/MathML')
    args = parser.parse_args()
    spec_path = Path(args.spec).resolve()
    spec = json.loads(spec_path.read_text(encoding='utf-8'))
    output_path = Path(args.output).resolve()
    # Figure paths resolve relative to spec, not output.
    greek = args.greek or (None if args.profile else 'times')    # an .eqp keeps its own Greek unless --greek is given
    profile = {'greek': greek, 'matrix_align': args.matrix_align, 'profile_json': args.profile}
    result = make_doc(spec, output_path, args.layout_preview, args.chrome_path,
                      args.write_equation_manifest, spec_path.parent, args.math_preview_scale,
                      profile, args.display_align, args.indent_twips)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
