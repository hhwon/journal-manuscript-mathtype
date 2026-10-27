#!/usr/bin/env python3
"""Generate DOCX with MathType OLE from self-owned MTEF and CFB code.

Uses style/metadata from a real MathType source object to reduce unknown
compatibility variables. Newly synthesized mathematical records come entirely
from mtef_subset.py. Images are independent PNG previews only.
"""
from __future__ import annotations
import argparse
from copy import deepcopy
from dataclasses import dataclass
from io import BytesIO
import json
from pathlib import Path
import hashlib
import struct
from zipfile import ZipFile, ZIP_DEFLATED

from lxml import etree
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

from ole_native import CompoundFile, equation_native_summary
from ole_cfb_writer import repack_stream
from ole_editor import FormulaDocument, W, R, O, V, PR, CT, NS
from mtef_subset import make_mtef, validate_subset_mtef
from mtef_explicit import make_mtef_explicit, validate_explicit
from native_preview import preview_png, preview_font_settings
from mtef_font_audit import read_font_profile

OLE_REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships/oleObject'
IMG_REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships/image'
FORMULAS=[
 ('A1','Basic Latin',r'E=kT','synthesized'),
 ('A2','Greek variables',r'\sigma=E\varepsilon','synthesized'),
 ('A3','Tensor subscripts',r'\sigma_{ij}=C_{ijkl}\varepsilon_{kl}','synthesized'),
 ('A4','Higher-gradient term',r'\sigma_{ij}=C_{ijkl}\varepsilon_{kl}+D_{ijklmn}\varepsilon_{kl,mn}','synthesized'),
 ('A5','Superscript',r'E=mc^{2}','synthesized'),
 ('A6','Fraction',r'\frac{1}{2}','synthesized'),
 ('A7','Nested radical',r'\sqrt{x^{2}+y^{2}}','synthesized'),
 ('B1','One-symbol MTEF edit',r'\tau=E\varepsilon','patched'),
 ('C1','Original MathType control',None,'original'),
]



def _make_style(doc):
    s=doc.sections[0]
    s.page_width=Inches(8.5);s.page_height=Inches(11)
    s.top_margin=s.bottom_margin=Inches(.82)
    s.left_margin=s.right_margin=Inches(.9)
    styles=doc.styles
    normal=styles['Normal']
    normal.font.name='Times New Roman'; normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(6)
    normal.paragraph_format.line_spacing=1.2
    for name, size in [('Title',17),('Heading 1',12.5),('Heading 2',11)]:
        style=styles[name];style.font.name='Times New Roman';style.font.size=Pt(size)
        style.font.bold=True
        style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.space_before=Pt(10)
        style.paragraph_format.space_after=Pt(5)
    title_style=styles['Title']
    for pb in title_style._element.xpath('./w:pPr/w:pBdr'):
        pb.getparent().remove(pb)
    if 'Equation ID' not in styles:
        style=styles.add_style('Equation ID',WD_STYLE_TYPE.PARAGRAPH)
        style.font.name='Times New Roman'; style.font.size=Pt(9)
        style.paragraph_format.space_after=Pt(1)
    footer=s.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('Linux MathType native OLE compatibility test  |  Windows validation pending').font.size=Pt(8)


def _make_table(doc, label, desc, expected, preview_status):
    desc_p=doc.add_paragraph(style='Equation ID')
    desc_p.paragraph_format.keep_with_next=True
    rr=desc_p.add_run(f'{label}  {desc}  |  {preview_status}')
    rr.bold=True
    table=doc.add_table(rows=1,cols=3)
    table.autofit=False
    table.columns[0].width=Inches(.4)
    table.columns[1].width=Inches(5.85)
    table.columns[2].width=Inches(.65)
    row=table.rows[0]
    for cell in row.cells:cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    row.cells[0].width=Inches(.4)
    row.cells[1].width=Inches(5.85)
    row.cells[2].width=Inches(.65)
    rp=row.cells[1].paragraphs[0]
    rp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    row.cells[2].paragraphs[0].alignment=WD_ALIGN_PARAGRAPH.RIGHT
    row.cells[2].paragraphs[0].add_run(f'({label})').font.size=Pt(9)
    if expected:
        p=doc.add_paragraph()
        p.paragraph_format.space_after=Pt(8)
        r=p.add_run('Expected LaTeX:  '+expected)
        r.font.name='Consolas';r.font.size=Pt(8)
    return rp


def _patch_one_char(native_ole, orig_code: int, new_code: int, orig_encoded:int, new_encoded:int):
    cf=CompoundFile(native_ole)
    native=cf.read('Equation Native')
    mt=native[28:]
    needle=b'\x02\x04\x84'+struct.pack('<H',orig_code)+bytes([orig_encoded])
    replacement=b'\x02\x04\x84'+struct.pack('<H',new_code)+bytes([new_encoded])
    if mt.count(needle)!=1:
        raise ValueError(f'Symbol must be uniquely present, found {mt.count(needle)}')
    edited=mt.replace(needle,replacement,1)
    if len(edited)!=len(mt):raise AssertionError('One-character edit changed MTEF length')
    ole=repack_stream(native_ole,'Equation Native',native[:28]+edited)
    cf2=CompoundFile(ole)
    assert cf2.read('Equation Native')[28:] == edited
    for k in cf.entries:
        if k!='Equation Native':assert cf.read(k)==cf2.read(k)
    return ole


def _serialize(root):
    return etree.tostring(root,xml_declaration=True,encoding='UTF-8',standalone=True)


def build(source: Path, output: Path, formulas=None):
    rows=list(FORMULAS if formulas is None else formulas)
    if not rows or len({r[0] for r in rows})!=len(rows):
        raise ValueError("Formula IDs must be nonempty and unique")
    source_doc=FormulaDocument(source)
    seed=source_doc.by_index(1)
    if seed.prog_id!='Equation.DSMT4' or seed.mtef_version!=5:
        raise ValueError('Seed must be genuine Equation.DSMT4 with MTEF v5')
    with ZipFile(source) as z:
        original_ole=z.read(seed.object_part)
        mtef_seed=CompoundFile(original_ole).read('Equation Native')[28:]
        font_profile=read_font_profile(mtef_seed)
        font_styles={v['name']:v for v in font_profile['styles']}
        if font_styles['Variable']['font'] != 'Times New Roman' or not font_styles['Variable']['italic']:
            raise ValueError('Unexpected source MTEF variable style; refuse silent font mismatch')
        if font_styles['Number']['font'] != 'Times New Roman' or font_styles['Number']['italic']:
            raise ValueError('Unexpected source MTEF number style; refuse silent font mismatch')
        control=source_doc.by_index(2)
        control_ole=z.read(control.object_part)
        control_preview=z.read(control.preview_part)
        donor_xml=deepcopy(source_doc.objects[0].getparent())
        control_xml=deepcopy(source_doc.objects[1].getparent())
    # Calculate all mathematical OLE contents before touching DOCX XML.
    payloads=[]
    synthesized={}
    for label,desc,latex,origin in rows:
        validation=None
        if origin in ('synthesized','synthesized-explicit'):
            if origin=='synthesized':
                mtef=make_mtef(latex,mtef_seed)
                validation=validate_subset_mtef(mtef)
            else:
                mtef=make_mtef_explicit(latex,mtef_seed)
                validation=validate_explicit(mtef)
            newheader=bytearray(CompoundFile(original_ole).read('Equation Native')[:28])
            struct.pack_into('<I',newheader,8,len(mtef))
            ole=repack_stream(original_ole,'Equation Native',bytes(newheader)+mtef)
            prv,w,h=preview_png(latex)
            synthesized[label]=ole
            image_ext='.png'
        elif origin=='patched':
            ole=_patch_one_char(synthesized['A2'],0x03C3,0x03C4,ord('s'),ord('t'))
            prv,w,h=preview_png(latex)
            image_ext='.png'
        else:
            ole=control_ole
            prv=control_preview
            image_ext=Path(control.preview_part).suffix.lower()
            w=h=None
        if equation_native_summary(ole)['mtef_version']!=5:
            raise AssertionError('Not MTEF v5')
        embedded_font_profile=read_font_profile(CompoundFile(ole).read('Equation Native')[28:])
        for name in ('fonts','styles','sizes','full_size'):
            if embedded_font_profile[name] != font_profile[name]:
                raise AssertionError('Native MTEF font profile diverged from MathType 7 source: '+name)
        if origin=='synthesized-explicit' and len(embedded_font_profile['explicit_font_styles'])!=5:
            raise AssertionError('Expected 5 MTEF explicit-font definitions')
        payloads.append((label,desc,latex,origin,ole,prv,w,h,image_ext))

    doc=Document();_make_style(doc)
    title=doc.add_paragraph(style='Title')
    title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    title.add_run('MathType Native OLE Compatibility Test')
    sub=doc.add_paragraph()
    sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run('Self-owned MTEF/CFB writer | Font semantics (Text / Function / Variable / Greek / Vector / Number)').italic=True
    doc.add_heading('Purpose and procedure',level=1)
    doc.add_paragraph('This file contains real Equation.DSMT4 MathType OLE parts, not Word OMML equations. Its editable MTEF styles specify 12 pt Times New Roman for variables/numbers/functions, and Symbol for Greek and operators. The independent PNG previews are rendered on Linux with Times-compatible serif typography. Windows MathType may regenerate the previews when the equations are edited.')
    for line in [
        '1. Open in Windows desktop Word with MathType 7 installed (prefer a disposable copy).',
        '2. Double-click the centre of each displayed equation. Confirm the MathType editor opens and the structure matches the expected LaTeX below it.',
        '3. Compare the unedited Word preview with the formula inside MathType. Then edit a symbol, save, close Word and reopen. Report whether either view has an incorrect font or size.',
    ]:
        p=doc.add_paragraph(line)
        p.paragraph_format.left_indent=Inches(.15)
    doc.add_heading('A. Newly encoded formulas',level=1)
    containers={}
    for label,desc,latex,origin,*_ in payloads:
        if label=='B1':doc.add_heading('B. Direct modification of MTEF character data',level=1)
        if label=='C1':doc.add_heading('C. Unchanged original MathType control',level=1)
        if label=='B1':
            doc.add_paragraph('B1 modifies one Greek CHAR record of A2 (sigma to tau) inside an Equation Native stream. The other OLE streams and formula structure are preserved by a CFB rebuild.')
        if label=='C1':
            doc.add_paragraph('C1 was copied unchanged from the supplied reference manuscript. It is included to distinguish an OLE packaging problem from a newly synthesized MTEF problem.')
        containers[label]=_make_table(doc,label,desc,latex,'new MTEF (explicit TNR)' if origin=='synthesized-explicit' else ('new MTEF (semantic)' if origin=='synthesized' else ('edited MTEF' if origin=='patched' else 'original control')))
    doc.add_page_break()
    doc.add_heading('Feedback checklist',level=1)
    tbl=doc.add_table(rows=1,cols=4);tbl.style='Table Grid'
    for cell,text in zip(tbl.rows[0].cells,['Formula','Visible','Opens MathType','Editable after reopen']):cell.text=text
    for label,*_ in payloads:
        row=tbl.add_row()
        row.cells[0].text=label
        for c in row.cells[1:]:c.text='Pass / Fail'
    doc.add_paragraph('If any equation fails, please report the formula ID and a screenshot or exact error text. A failed equation is not proof the OLE object is not present; it distinguishes an application compatibility issue from structural packaging.')
    temp=output.with_name('_ole_prebase.docx')
    output.parent.mkdir(parents=True,exist_ok=True)
    doc.save(temp)
    image_parts={}
    ole_parts={}
    meta=[]
    with ZipFile(temp) as z:
        root=etree.fromstring(z.read('word/document.xml'))
        rels=etree.fromstring(z.read('word/_rels/document.xml.rels'))
        types=etree.fromstring(z.read('[Content_Types].xml'))
        max_id=max((int(e.get('Id')[3:]) for e in rels if e.get('Id','').startswith('rId') and e.get('Id','')[3:].isdigit()),default=50)
        for ix,(label,desc,latex,origin,ole,prv,w,h,ext) in enumerate(payloads,1):
            max_id+=1; image_rid=f'rId{max_id}'
            max_id+=1; ole_rid=f'rId{max_id}'
            p=containers[label]
            # Paragraph/runs are proxies sharing same underlying lxml nodes as
            # the document saved to 'temp', but do not reserialize via docx API.
            # Find unique label row by position (the right table cell contains
            # (ID)); retrieve the corresponding centre paragraph in XML.
            tables=root.xpath('.//w:body/w:tbl',namespaces=NS)
            matching=[]
            for t in tables:
                txt=''.join(t.xpath('.//w:t/text()',namespaces=NS))
                if '('+label+')' in txt:
                    matching.append(t)
            if len(matching)!=1:raise ValueError(f'Ambiguous formula table {label}')
            center=matching[0].xpath('./w:tr/w:tc[2]/w:p',namespaces=NS)[0]
            r=etree.SubElement(center,'{%s}r'%W)
            obj=deepcopy(control_xml if origin=='original' else donor_xml)
            obj.attrib.pop('{http://schemas.microsoft.com/office/word/2010/wordml}anchorId',None)
            shape=obj.xpath('.//v:shape',namespaces=NS)[0]
            ole_object=obj.xpath('./o:OLEObject',namespaces=NS)[0]
            imagedata=obj.xpath('.//v:imagedata',namespaces=NS)[0]
            sid=f'_x0000_i{3100+ix}'
            shape.set('id',sid)
            ole_object.set('ShapeID',sid)
            ole_object.set('ObjectID',f'_249000{ix:04d}')
            ole_object.set('{%s}id'%R,ole_rid)
            imagedata.set('{%s}id'%R,image_rid)
            if w is not None:
                shape.set('style',f'width:{w:.2f}pt;height:{h:.2f}pt')
                obj.set('{%s}dxaOrig'%W,str(int(w*20)))
                obj.set('{%s}dyaOrig'%W,str(int(h*20)))
            r.append(obj)
            imagepath=f'word/media/native-mtef-{ix}{ext}'
            olepath=f'word/embeddings/native-mtef-{ix}.bin'
            image_parts[imagepath]=prv
            ole_parts[olepath]=ole
            etree.SubElement(rels,'{%s}Relationship'%PR,Id=image_rid,Type=IMG_REL,
                             Target=f'media/native-mtef-{ix}{ext}')
            etree.SubElement(rels,'{%s}Relationship'%PR,Id=ole_rid,Type=OLE_REL,
                             Target=f'embeddings/native-mtef-{ix}.bin')
            meta.append({'test_id':label,'origin':origin,'latex':latex,
                         'object_part':olepath,'preview_part':imagepath,
                         'sha256':hashlib.sha256(ole).hexdigest(),
                         'mtef':equation_native_summary(ole),
                         'typefaces': (validate_subset_mtef(CompoundFile(ole).read('Equation Native')[28:])['style_counts']
                                       if origin=='synthesized' else (validate_explicit(CompoundFile(ole).read('Equation Native')[28:])['explicit_typefaces']
                                         if origin=='synthesized-explicit' else None))})
        defaults={e.get('Extension') for e in types if e.tag.endswith('Default')}
        for extension,ctype in [('bin','application/vnd.openxmlformats-officedocument.oleObject'),
                                ('png','image/png'),('wmf','image/x-wmf')]:
            if extension not in defaults:
                etree.SubElement(types,'{%s}Default'%CT,Extension=extension,ContentType=ctype)
        changed={'word/document.xml':_serialize(root),
                 'word/_rels/document.xml.rels':_serialize(rels),
                 '[Content_Types].xml':_serialize(types)}
        with ZipFile(output,'w',ZIP_DEFLATED) as out:
            for item in z.infolist():
                out.writestr(item,changed.get(item.filename,z.read(item.filename)))
            for name,data in image_parts.items():out.writestr(name,data)
            for name,data in ole_parts.items():out.writestr(name,data)
    temp.unlink(missing_ok=True)
    verified=FormulaDocument(output)
    if len(verified.formulas)!=len(rows):raise AssertionError('Formula inventory mismatch')
    for i in range(len(rows)):
        f=verified.by_index(i+1)
        if f.object_sha256!=meta[i]['sha256']:
            raise AssertionError('Embedded MTEF hash mismatch')
    if any(r[3]=='patched' for r in rows):
        with ZipFile(output) as z:
            patch_index=next(i for i,r in enumerate(rows,1) if r[3]=='patched')
            if 'A2' not in [r[0] for r in rows]:
                raise AssertionError('Patched control requires A2 reference')
            idx_a2=next(i for i,r in enumerate(rows,1) if r[0]=='A2')
            b=z.read(verified.by_index(idx_a2).object_part)
            edit=z.read(verified.by_index(patch_index).object_part)
            va=CompoundFile(b).read('Equation Native')[28:]
            ve=CompoundFile(edit).read('Equation Native')[28:]
            assert len(va)==len(ve)
            diff=[i for i in range(len(va)) if va[i]!=ve[i]]
            assert len(diff)==2,(len(diff),diff)
            assert CompoundFile(edit).summary()==CompoundFile(b).summary()
    manifest=output.with_suffix('.manifest.json')
    manifest.write_text(json.dumps({'filename':str(output),'seed':str(source),
                                    'preview_font':preview_font_settings()[0],
                                    'mtef_font_profile':font_profile,
                                    'mathtype_editing_verified_on_windows':False,
                                    'tests':meta},indent=2,ensure_ascii=False),encoding='utf-8')
    print(f'SUCCESS generated {output}, MathType objects {len(verified.formulas)}')
    print(f'MANIFEST {manifest}')
    return output


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--source',required=True,help='A genuine DOCX with Equation.DSMT4 MathType v5 OLE')
    p.add_argument('-o','--output',required=True)
    a=p.parse_args()
    build(Path(a.source),Path(a.output))

if __name__=='__main__':main()
