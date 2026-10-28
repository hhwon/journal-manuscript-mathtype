#!/usr/bin/env python3
"""Create a clean Word template based on the second provided manuscript.

Does not copy draft text, broken equation fields, or author's OLE objects.
"""
from pathlib import Path
import argparse

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _field(paragraph, instr):
    r = paragraph.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin'); r._r.append(begin)
    r = paragraph.add_run()
    tag = OxmlElement('w:instrText'); tag.set(qn('xml:space'), 'preserve'); tag.text = instr; r._r.append(tag)
    r = paragraph.add_run()
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate'); r._r.append(sep)
    r = paragraph.add_run('1')
    r = paragraph.add_run()
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end'); r._r.append(end)


def _style(doc, name, *, size=12, font='Times New Roman', bold=False, italic=False,
           align=None, line=1.5, first_cm=0.0, before=0, after=0, keep_next=False):
    styles = doc.styles
    s = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = styles['Normal'] if name != 'Normal' else None
    s.font.name = font
    s.font.size = Pt(size)
    s.font.bold = bold
    s.font.italic = italic
    s.font.color.rgb = None
    s.paragraph_format.alignment = align
    s.paragraph_format.line_spacing = line
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    # every style sets its own first-line indent (0 unless given): otherwise headings, figures, captions, equations and
    # table cells inherit the 0.85 cm body indent from Normal
    s.paragraph_format.first_line_indent = Cm(first_cm)
    s.paragraph_format.keep_with_next = keep_next
    rpr = s.element.get_or_add_rPr()
    rf = rpr.rFonts
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ['ascii','hAnsi','cs']:
        rf.set(qn('w:'+a), font)
    rf.set(qn('w:eastAsia'), 'SimSun')
    return s


def _heading(doc, name, level, **kw):
    """A heading style that Word lists in the Navigation Pane and in a table of contents: the built-in
    'Heading n' style (outline level n-1) with the manuscript formatting; theme fonts and colours removed."""
    s = _style(doc, name, **kw)
    rpr = s.element.get_or_add_rPr()
    rf = rpr.rFonts
    for a in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        if rf is not None and rf.get(qn('w:' + a)) is not None:
            del rf.attrib[qn('w:' + a)]
    for tag in ('w:color', 'w:szCs'):
        for el in rpr.findall(qn(tag)):
            rpr.remove(el)
    ppr = s.element.get_or_add_pPr()
    for el in ppr.findall(qn('w:outlineLvl')):
        ppr.remove(el)
    lvl = OxmlElement('w:outlineLvl'); lvl.set(qn('w:val'), str(level - 1))
    ppr.append(lvl)
    link = s.element.find(qn('w:link'))           # the linked character style must not bring back theme fonts/colours
    if link is not None and link.get(qn('w:val')) in doc.styles:
        crpr = doc.styles[link.get(qn('w:val'))].element.get_or_add_rPr()
        for el in list(crpr):
            if el.tag in (qn('w:rFonts'), qn('w:color'), qn('w:sz'), qn('w:szCs')):
                crpr.remove(el)
    return s


def create_template(output):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.left_margin = sec.right_margin = Inches(1)
    sec.top_margin = sec.bottom_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(.5)

    _style(doc, 'Normal', size=12, line=1.5, first_cm=0.85, after=0,
           align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    _style(doc, 'MS Title', size=16, bold=True, line=1.5, after=6,
           align=WD_ALIGN_PARAGRAPH.CENTER)
    _style(doc, 'MS Authors', size=12, line=1.5,
           align=WD_ALIGN_PARAGRAPH.CENTER)
    _style(doc, 'MS Affiliation', size=12, line=1.15, after=4,
           align=WD_ALIGN_PARAGRAPH.CENTER)
    _style(doc, 'MS Corresponding', size=10, line=1.15, after=12,
           align=WD_ALIGN_PARAGRAPH.CENTER)
    _style(doc, 'MS Abstract Heading', size=12, bold=True,
           line=1.5, before=18, after=6, align=WD_ALIGN_PARAGRAPH.CENTER,
           keep_next=True)
    _style(doc, 'MS Abstract', size=12, line=1.5, first_cm=.85, after=0,
           align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    _style(doc, 'MS Keywords', size=12, line=1.5, first_cm=.85, before=6, after=12,
           align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    # section headings use Word's built-in heading styles, so that they carry outline levels (Navigation Pane, TOC);
    # 'MS Heading 1/2' are kept as identical aliases for older builders
    for name, level in (('Heading 1', 1), ('MS Heading 1', 1)):
        _heading(doc, name, level, font='Arial', size=14, bold=True,
                 line=1.5, before=18, after=6, keep_next=True,
                 align=WD_ALIGN_PARAGRAPH.LEFT)
    for name, level in (('Heading 2', 2), ('MS Heading 2', 2)):
        _heading(doc, name, level, size=12, italic=True,
                 line=1.5, before=0, after=0, keep_next=True,
                 align=WD_ALIGN_PARAGRAPH.LEFT)
    _style(doc, 'MS Body', size=12, line=1.5, first_cm=.85, after=0,
           align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    _style(doc, 'MS Equation', size=12, line=1.2, before=2, after=2,
           align=WD_ALIGN_PARAGRAPH.CENTER)
    _style(doc, 'MS Figure', size=10, line=1.0, before=4, after=0,
           align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
    _style(doc, 'MS Figure Caption', size=12, line=1.5, before=6, after=12,
           align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    _style(doc, 'MS Table Caption', size=12, line=1.5, before=12, after=6,
           align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep_next=True)
    _style(doc, 'MS Table Cell', size=10, line=1.0,
           align=WD_ALIGN_PARAGRAPH.CENTER)
    _style(doc, 'MS References', size=11, line=1.0, first_cm=-1.0, after=5,
           align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    doc.styles['MS References'].paragraph_format.left_indent = Cm(1.0)
    doc.styles['MS References'].paragraph_format.tab_stops.add_tab_stop(Cm(1.0))
    _style(doc, 'MS Note', size=10, line=1.2, after=5,
           align=WD_ALIGN_PARAGRAPH.LEFT)

    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot.style = doc.styles['MS Note']
    _field(foot, ' PAGE ')

    doc.core_properties.title = 'Research manuscript template'
    doc.core_properties.subject = 'Journal manuscript style (MathType OLE ready)'
    doc.core_properties.author = ''
    doc.core_properties.comments = ''
    out = Path(output)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(f'Created {out}')
    return out


if __name__ == '__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('output', nargs='?', default=str(Path(__file__).resolve().parent.parent/'assets'/'manuscript-template.docx'))
    create_template(parser.parse_args().output)
